from __future__ import annotations

import math
import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from rb.cache import ArtifactCache
from rb.net import http_get
from rb.util import write_json_atomic, write_text_atomic


def parse_household_workbook(body: bytes, cfg: dict) -> tuple[list[tuple[int, float]], dict[int, str]]:
    """Read a published annual column, resolving only documented duplicate years."""
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    try:
        sheet = workbook[cfg["sheet"]]
        for cell, expected in cfg["expected_cells"].items():
            actual = " ".join(str(sheet[cell].value).split())
            if actual != expected:
                raise ValueError(f"Unexpected workbook header at {cell}: {actual!r}")
        selected: dict[int, float] = {}
        labels: dict[int, str] = {}
        seen: dict[int, list[str]] = {}
        preferred = {int(k): str(v) for k, v in cfg.get("preferred_year_labels", {}).items()}
        column = int(cfg["value_column"])
        scale = float(cfg.get("scale", 1))
        for row in sheet.iter_rows(min_row=int(cfg["first_data_row"]), values_only=True):
            label = str(row[0]).strip()
            match = re.fullmatch(r"(\d{4})(?:\s+\d+(?:,\s*\d+)*|[ab])?\*?", label)
            if not match:
                continue
            year = int(match[1])
            seen.setdefault(year, []).append(label)
            if year in preferred and label != preferred[year]:
                continue
            if year in selected:
                raise ValueError(f"Unresolved duplicate year: {year}")
            value = float(row[column - 1]) * scale
            if not math.isfinite(value) or not cfg["minimum"] <= value <= cfg["maximum"]:
                raise ValueError(f"Invalid value for {year}: {value}")
            selected[year] = value
            labels[year] = label
        for year in seen:
            if year not in selected:
                raise ValueError(f"Preferred row missing for {year}: {preferred.get(year)}")
        if not selected:
            raise ValueError("Workbook produced no observations")
        years = sorted(selected)
        if years != list(range(int(cfg["start_year"]), int(cfg["end_year"]) + 1)):
            raise ValueError("Workbook annual coverage differs from the registered release")
        return sorted(selected.items()), labels
    finally:
        workbook.close()


def ingest_household_series(
    *, source_name: str, series_key: str, series_cfg: dict, source_cfg: dict, refresh: bool,
) -> None:
    cache = ArtifactCache()
    raw_dir = cache.artifact_dir("household", source_name)
    artifact = cache.latest(raw_dir, suffix="xlsx")
    if refresh or artifact is None:
        url = source_cfg["url"]
        status, headers, body = http_get(url)
        artifact = cache.write(raw_dir, data=body, suffix="xlsx", meta={
            "url": url, "status": status, "headers": headers,
        })
    rows, labels = parse_household_workbook(artifact.path.read_bytes(), series_cfg["parse"])
    output = Path("data/derived/household") / f"{series_key}.csv"
    output.parent.mkdir(parents=True, exist_ok=True)
    write_text_atomic(output, "date,value\n" + "".join(f"{y}-01-01,{v:.12g}\n" for y, v in rows))
    write_json_atomic(output.with_suffix(".meta.json"), {
        "url": source_cfg["url"], "sha256": artifact.sha256,
        "parse": series_cfg["parse"], "selected_year_labels": labels,
        "units": series_cfg["units"], "frequency": "A", "observations": len(rows),
    })
